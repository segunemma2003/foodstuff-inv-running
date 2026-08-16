"""Invoice domain services."""

from datetime import date
from io import BytesIO
from typing import List, Optional
import uuid
import os

from fastapi import HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

import models
import schemas
from utils import audit
from utils.make_integration import send_document_to_make_from_s3
from utils.number_gen import next_invoice_number
from utils.pricing import get_active_rules, calculate_item_price
from decimal import Decimal
from utils.pdf_generator import generate_cost_of_sales_pdf
from utils.pdf_generator import generate_invoice_pdf as gen_pdf
from utils.queue_events import log_queue_event
from utils.s3 import delete_object, download_bytes, upload_bytes
from utils.tasks import (
    generate_invoice_pdf_task,
    process_invoice_bulk_task,
    send_invoice_to_recipients_task,
)


INVOICE_PRIMARY_RECIPIENT = os.getenv(
    "INVOICE_PRIMARY_RECIPIENT",
    os.getenv("INVOICE_PRIMARY_RECIPIENT_EMAIL", "foodstuffstoreinvoices@gmail.com"),
)


def list_convertible_quotations(db: Session, customer_id: Optional[int] = None):
    convertible_quotation_query = (
        db.query(models.Quotation)
        .outerjoin(models.Invoice, models.Invoice.quotation_id == models.Quotation.id)
        .filter(
            models.Quotation.status == models.QuotationStatus.approved,
            models.Invoice.id == None,
        )
    )
    if customer_id:
        convertible_quotation_query = convertible_quotation_query.filter(
            models.Quotation.customer_id == customer_id
        )
    return convertible_quotation_query.order_by(models.Quotation.approved_at.desc()).all()


def download_invoice_template():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice Import"

    headers = [
        ("invoice_number", 20),
        ("customer_name", 28),
        ("invoice_date", 16),
        ("due_date", 16),
        ("payment_term", 18),
        ("delivery_type", 16),
        ("product_name", 30),
        ("qty", 10),
        ("unit_price", 14),
        ("notes", 30),
    ]

    green_fill = PatternFill("solid", fgColor="1E8449")
    alt_fill = PatternFill("solid", fgColor="EAF4EE")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    body_font = Font(size=10)
    thin_side = Side(style="thin", color="CCCCCC")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for col, (name, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = green_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22

    sample = [
        ("INV-2026-0001", "GENESIS GROUP", date(2026, 1, 15), date(2026, 2, 14), "net_30", "delivery", "Rice 50kg", 10, 85000, ""),
        ("INV-2026-0001", "GENESIS GROUP", date(2026, 1, 15), date(2026, 2, 14), "net_30", "delivery", "Beans 50kg", 5, 45000, ""),
        ("", "ACME Corp", date(2026, 1, 16), "", "cash", "pickup", "Rice 50kg", 2, 85000, "Urgent order"),
    ]

    date_fmt = "YYYY-MM-DD"
    for r, row in enumerate(sample, 2):
        fill = alt_fill if r % 2 == 0 else None
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = body_font
            cell.border = thin_border
            if fill:
                cell.fill = fill
            if isinstance(val, date):
                cell.number_format = date_fmt
                cell.alignment = Alignment(horizontal="center")
            elif c in (8, 9):
                cell.alignment = Alignment(horizontal="right")
                if c == 9:
                    cell.number_format = "#,##0"

    ws.freeze_panes = "A2"

    ns = wb.create_sheet("How to Fill")
    ns.column_dimensions["A"].width = 18
    ns.column_dimensions["B"].width = 85
    ns["A1"] = "Column"
    ns["B1"] = "Instructions"
    ns["A1"].font = Font(bold=True, color="FFFFFF")
    ns["B1"].font = Font(bold=True, color="FFFFFF")
    ns["A1"].fill = green_fill
    ns["B1"].fill = green_fill

    instructions = [
        ("invoice_number", "Optional. Leave blank to auto-generate. Use the SAME value on multiple rows to group them into one invoice."),
        ("customer_name", "REQUIRED. Must match a customer name in the system exactly (case-insensitive)."),
        ("invoice_date", "REQUIRED. Date format: YYYY-MM-DD  (e.g. 2026-01-15)"),
        ("due_date", "Optional. Date format: YYYY-MM-DD"),
        ("payment_term", "Optional. Allowed values: cash  immediate  net_7  net_14  net_30  net_45  net_60  net_90.  Default: cash"),
        ("delivery_type", "Optional. Allowed values: delivery  pickup.  Default: pickup"),
        ("product_name", "REQUIRED. Must match a product name in the system exactly (case-insensitive)."),
        ("qty", "REQUIRED. Quantity sold. Must be greater than 0."),
        ("unit_price", "REQUIRED. Selling price per unit — numbers only, no currency symbol (e.g. 85000)."),
        ("notes", "Optional. Any internal note for this invoice."),
    ]
    for i, (col, desc) in enumerate(instructions, 2):
        a = ns.cell(row=i, column=1, value=col)
        b = ns.cell(row=i, column=2, value=desc)
        a.font = Font(bold=True, size=10)
        b.font = Font(size=10)
        if i % 2 == 0:
            a.fill = alt_fill
            b.fill = alt_fill

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="invoice_import_template.xlsx"'},
    )


async def bulk_upload_invoices(file: UploadFile, current_user: models.User) -> schemas.JobEnqueuedResponse:
    content = await file.read()
    s3_key = f"uploads/invoices_{uuid.uuid4()}.xlsx"
    upload_bytes(s3_key, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    task = process_invoice_bulk_task.delay(s3_key, current_user.id)
    return schemas.JobEnqueuedResponse(
        task_id=task.id,
        message=f"Invoice import queued. Poll /api/v1/jobs/{task.id} for result.",
    )


def list_invoices(
    db: Session,
    *,
    skip: int = 0,
    limit: int = 50,
    customer_id: Optional[int] = None,
    status: Optional[str] = None,
    created_by: Optional[int] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    payment_term: Optional[str] = None,
    delivery_type: Optional[str] = None,
) -> schemas.InvoiceListPage:
    skip = max(skip, 0)
    limit = min(max(limit, 1), 200)

    invoice_query = db.query(models.Invoice).options(joinedload(models.Invoice.customer))
    if customer_id:
        invoice_query = invoice_query.filter(models.Invoice.customer_id == customer_id)
    if status:
        invoice_query = invoice_query.filter(models.Invoice.status == status)
    if created_by:
        invoice_query = invoice_query.filter(models.Invoice.created_by == created_by)
    if date_from:
        invoice_query = invoice_query.filter(models.Invoice.invoice_date >= date_from)
    if date_to:
        invoice_query = invoice_query.filter(models.Invoice.invoice_date <= date_to)
    if payment_term:
        invoice_query = invoice_query.filter(models.Invoice.payment_term == payment_term)
    if delivery_type:
        invoice_query = invoice_query.filter(models.Invoice.delivery_type == delivery_type)

    total = invoice_query.order_by(None).count()
    items = (
        invoice_query.order_by(models.Invoice.created_at.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return schemas.InvoiceListPage(total=total, skip=skip, limit=limit, items=items)


def _delete_invoice_record(db: Session, invoice_id: int, current_user: models.User) -> schemas.MessageResponse:
    inv = db.query(models.Invoice).filter(models.Invoice.id == invoice_id).first()
    if not inv:
        raise HTTPException(404, "Invoice not found")

    quotation_id = inv.quotation_id
    inv_no = inv.invoice_number

    for payment in list(inv.payments):
        db.delete(payment)

    if inv.custom_pdf_s3_key:
        try:
            delete_object(inv.custom_pdf_s3_key)
        except Exception:
            pass

    audit.log(
        db,
        models.AuditAction.delete,
        models.AuditEntity.invoice,
        inv.id,
        current_user.id,
        description=f"Deleted invoice {inv_no}",
    )
    db.delete(inv)

    if quotation_id:
        qt = db.query(models.Quotation).filter(models.Quotation.id == quotation_id).first()
        if qt and qt.status == models.QuotationStatus.converted:
            qt.status = models.QuotationStatus.approved

    return schemas.MessageResponse(message=f"Invoice {inv_no} deleted")


def bulk_delete_invoices(
    db: Session, body: schemas.BulkIdsRequest, current_user: models.User
) -> schemas.BulkDeleteResult:
    result = schemas.BulkDeleteResult()
    for iid in body.ids:
        try:
            _delete_invoice_record(db, iid, current_user)
            result.deleted += 1
        except HTTPException as he:
            result.failed.append(
                {"id": iid, "detail": he.detail if isinstance(he.detail, str) else str(he.detail)}
            )
    db.commit()
    return result



def _resolve_invoice_unit_price(item, *, cost_price: int, delivery_type: str, payment_term: str, rules: dict, wht_enabled: bool) -> tuple[int, dict]:
    """Return (unit_price, pricing breakdown). Override wins; else cost + markups [+ WHT]."""
    override = getattr(item, "unit_price_override", None)
    if override is None:
        override = getattr(item, "unit_price", None)
    pricing = calculate_item_price(
        Decimal(str(cost_price or 0)),
        delivery_type,
        payment_term,
        rules,
        wht_enabled=wht_enabled,
    )
    if override is not None:
        return int(override), pricing
    return int(pricing["unit_price"]), pricing


def create_invoice(db: Session, body: schemas.InvoiceCreate, current_user: models.User) -> models.Invoice:
    customer = db.query(models.Customer).filter(models.Customer.id == body.customer_id).first()
    if not customer:
        raise HTTPException(404, "Customer not found")
    if not body.items:
        raise HTTPException(400, "At least one line item is required")

    try:
        delivery_type = models.DeliveryType(body.delivery_type)
    except ValueError:
        raise HTTPException(400, f"Invalid delivery_type '{body.delivery_type}'")

    # Prefer admin-selected term; if omitted, use the customer's default payment term
    payment_term = (body.payment_term or "").strip() or None
    if not payment_term:
        payment_term = (customer.default_payment_term or "").strip() or "cash"

    apply_wht = bool(getattr(body, "apply_wht", False))
    rules = get_active_rules(db)
    line_items = []
    total_amount = 0

    for item in body.items:
        product = db.query(models.Product).filter(models.Product.id == item.product_id).first()
        if not product:
            raise HTTPException(404, f"Product {item.product_id} not found")
        if not product.is_active:
            raise HTTPException(
                400, f"Product '{product.product_name}' is deactivated and cannot be used"
            )

        latest_cost = (
            db.query(models.CostPrice)
            .filter(models.CostPrice.product_id == product.id)
            .order_by(models.CostPrice.effective_date.desc())
            .first()
        )
        cost_price = int(latest_cost.cost_price) if latest_cost else 0
        unit_price, pricing = _resolve_invoice_unit_price(
            item,
            cost_price=cost_price,
            delivery_type=body.delivery_type,
            payment_term=payment_term,
            rules=rules,
            wht_enabled=apply_wht,
        )
        quantity = float(item.quantity)
        line_total = int(round(unit_price * quantity))
        total_amount += line_total

        line_items.append(
            models.InvoiceItem(
                product_id=product.id,
                quantity=quantity,
                uom=item.uom or product.unit_of_measure,
                cost_price=cost_price,
                supply_markup_pct=Decimal(str(pricing["supply_markup_pct"])),
                supply_markup_amount=pricing["supply_markup_amount"],
                delivery_markup_pct=Decimal(str(pricing["delivery_markup_pct"])),
                delivery_markup_amount=pricing["delivery_markup_amount"],
                payment_term_markup_pct=Decimal(str(pricing["payment_term_markup_pct"])),
                payment_term_markup_amount=pricing["payment_term_markup_amount"],
                discount_pct=0,
                unit_price=unit_price,
                line_total=line_total,
            )
        )

    handling_charge_amount = int(round(total_amount * 5 / 100)) if body.apply_handling_charge else 0
    grand_total = total_amount + handling_charge_amount

    due_date = body.due_date
    if due_date is None and body.invoice_date:
        from datetime import timedelta
        _days = {
            "net_3": 3, "net_5": 5, "net_7": 7, "net_15": 15,
            "net_30": 30, "net_45": 45, "net_60": 60, "net_90": 90,
            "credit_3": 3, "credit_5": 5, "credit_7": 7,
            "credit_15": 15, "credit_15d": 15,
            "credit_30": 30, "credit_30d": 30,
            "credit_60": 60, "credit_60d": 60,
            "credit_90": 90, "credit_90d": 90,
        }.get((payment_term or "").lower())
        if _days:
            due_date = body.invoice_date + timedelta(days=_days)

    inv = models.Invoice(
        invoice_number=next_invoice_number(db),
        quotation_id=None,
        customer_id=customer.id,
        invoice_date=body.invoice_date,
        due_date=due_date,
        payment_term=payment_term,
        delivery_type=delivery_type,
        notes=body.notes,
        handling_charge_enabled=body.apply_handling_charge,
        handling_charge_amount=handling_charge_amount,
        wht_enabled=apply_wht,
        total_amount=grand_total,
        amount_paid=0,
        created_by=current_user.id,
        items=line_items,
    )
    db.add(inv)
    db.commit()
    db.refresh(inv)
    audit.log(
        db,
        models.AuditAction.create,
        models.AuditEntity.invoice,
        inv.id,
        current_user.id,
        description=f"Created invoice {inv.invoice_number} directly",
    )
    return inv


def update_invoice(db: Session, invoice_id: int, body: schemas.InvoiceUpdate, current_user: models.User) -> models.Invoice:
    inv = get_invoice(db, invoice_id)
    if inv.status == models.InvoiceStatus.cancelled:
        raise HTTPException(400, "Cannot edit a cancelled invoice")

    if body.customer_id is not None:
        customer = db.query(models.Customer).filter(models.Customer.id == body.customer_id).first()
        if not customer:
            raise HTTPException(404, "Customer not found")
        inv.customer_id = body.customer_id

    if body.invoice_date is not None:
        inv.invoice_date = body.invoice_date
    if body.due_date is not None:
        inv.due_date = body.due_date
    if body.payment_term is not None:
        inv.payment_term = body.payment_term
    if body.delivery_type is not None:
        try:
            inv.delivery_type = models.DeliveryType(body.delivery_type)
        except ValueError:
            raise HTTPException(400, f"Invalid delivery_type '{body.delivery_type}'")
    if body.notes is not None:
        inv.notes = body.notes
    wht_changed = body.apply_wht is not None and bool(body.apply_wht) != bool(inv.wht_enabled)
    if body.apply_wht is not None:
        inv.wht_enabled = bool(body.apply_wht)

    delivery_str = (
        inv.delivery_type.value if hasattr(inv.delivery_type, "value") else str(inv.delivery_type)
    )
    payment_str = inv.payment_term or "cash"
    wht_on = bool(inv.wht_enabled)

    if body.items is not None:
        for old_item in list(inv.items):
            db.delete(old_item)
        db.flush()

        rules = get_active_rules(db)
        line_items = []
        total_amount = 0
        for item in body.items:
            product = db.query(models.Product).filter(models.Product.id == item.product_id).first()
            if not product:
                raise HTTPException(404, f"Product {item.product_id} not found")
            if not product.is_active:
                raise HTTPException(400, f"Product '{product.product_name}' is deactivated and cannot be used")

            latest_cost = (
                db.query(models.CostPrice)
                .filter(models.CostPrice.product_id == product.id)
                .order_by(models.CostPrice.effective_date.desc())
                .first()
            )
            cost_price = int(latest_cost.cost_price) if latest_cost else 0
            # When WHT flag changes, ignore client unit prices and rebuild from cost + rules.
            # Otherwise honor override/unit_price when provided.
            if wht_changed:
                class _NoOverride:
                    unit_price = None
                    unit_price_override = None
                priced_from = _NoOverride()
            else:
                priced_from = item
            unit_price, pricing = _resolve_invoice_unit_price(
                priced_from,
                cost_price=cost_price,
                delivery_type=delivery_str,
                payment_term=payment_str,
                rules=rules,
                wht_enabled=wht_on,
            )
            # If client sent an explicit price and WHT did not change, keep it
            if not wht_changed:
                override = getattr(item, "unit_price_override", None)
                if override is None:
                    override = getattr(item, "unit_price", None)
                if override is not None:
                    unit_price = int(override)
            quantity = float(item.quantity)
            line_total = int(round(unit_price * quantity))
            total_amount += line_total

            line_items.append(
                models.InvoiceItem(
                    invoice_id=inv.id,
                    product_id=product.id,
                    quantity=quantity,
                    uom=item.uom or product.unit_of_measure,
                    cost_price=cost_price,
                    supply_markup_pct=Decimal(str(pricing["supply_markup_pct"])),
                    supply_markup_amount=pricing["supply_markup_amount"],
                    delivery_markup_pct=Decimal(str(pricing["delivery_markup_pct"])),
                    delivery_markup_amount=pricing["delivery_markup_amount"],
                    payment_term_markup_pct=Decimal(str(pricing["payment_term_markup_pct"])),
                    payment_term_markup_amount=pricing["payment_term_markup_amount"],
                    discount_pct=0,
                    unit_price=unit_price,
                    line_total=line_total,
                )
            )

        apply_handling = body.apply_handling_charge if body.apply_handling_charge is not None else inv.handling_charge_enabled
        handling_charge_amount = int(round(total_amount * 5 / 100)) if apply_handling else 0
        inv.items = line_items
        inv.handling_charge_enabled = apply_handling
        inv.handling_charge_amount = handling_charge_amount
        inv.total_amount = total_amount + handling_charge_amount
    elif wht_changed:
        # Reprice existing lines from cost so WHT is baked into unit prices
        rules = get_active_rules(db)
        total_amount = 0
        for ii in inv.items:
            pricing = calculate_item_price(
                Decimal(str(ii.cost_price or 0)),
                delivery_str,
                payment_str,
                rules,
                wht_enabled=wht_on,
            )
            ii.supply_markup_pct = Decimal(str(pricing["supply_markup_pct"]))
            ii.supply_markup_amount = pricing["supply_markup_amount"]
            ii.delivery_markup_pct = Decimal(str(pricing["delivery_markup_pct"]))
            ii.delivery_markup_amount = pricing["delivery_markup_amount"]
            ii.payment_term_markup_pct = Decimal(str(pricing["payment_term_markup_pct"]))
            ii.payment_term_markup_amount = pricing["payment_term_markup_amount"]
            ii.unit_price = int(pricing["unit_price"])
            ii.line_total = int(round(float(ii.quantity) * ii.unit_price))
            total_amount += ii.line_total
        apply_handling = body.apply_handling_charge if body.apply_handling_charge is not None else inv.handling_charge_enabled
        handling_charge_amount = int(round(total_amount * 5 / 100)) if apply_handling else 0
        inv.handling_charge_enabled = apply_handling
        inv.handling_charge_amount = handling_charge_amount
        inv.total_amount = total_amount + handling_charge_amount
    elif body.apply_handling_charge is not None:
        subtotal = sum(i.line_total for i in inv.items)
        handling_charge_amount = int(round(subtotal * 5 / 100)) if body.apply_handling_charge else 0
        inv.handling_charge_enabled = body.apply_handling_charge
        inv.handling_charge_amount = handling_charge_amount
        inv.total_amount = subtotal + handling_charge_amount

    if inv.quotation_id:
        qt = db.query(models.Quotation).filter(models.Quotation.id == inv.quotation_id).first()
        if qt:
            if body.delivery_type is not None:
                qt.delivery_type = inv.delivery_type
            if body.payment_term is not None:
                qt.payment_term = inv.payment_term
            if body.notes is not None:
                qt.notes = inv.notes
            if body.items is not None or body.apply_handling_charge is not None:
                qt.handling_charge_enabled = inv.handling_charge_enabled
                qt.handling_charge_amount = inv.handling_charge_amount
                qt.total_amount = inv.total_amount

    audit.log(
        db,
        models.AuditAction.update,
        models.AuditEntity.invoice,
        inv.id,
        current_user.id,
        description=f"Updated invoice {inv.invoice_number}",
    )
    db.commit()
    db.refresh(inv)
    return inv


def get_invoice(db: Session, invoice_id: int) -> models.Invoice:
    inv = db.query(models.Invoice).filter(models.Invoice.id == invoice_id).first()
    if not inv:
        raise HTTPException(404, "Invoice not found")
    return inv


def download_invoice_pdf(db: Session, invoice_id: int):
    inv = get_invoice(db, invoice_id)
    if inv.custom_pdf_s3_key:
        pdf_bytes = download_bytes(inv.custom_pdf_s3_key)
        return StreamingResponse(
            BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{inv.invoice_number}.pdf"'},
        )

    from utils.bank_accounts import resolve_bank_accounts_for_customer
    bank_accounts = resolve_bank_accounts_for_customer(db, inv.customer_id)
    paystack_payment = (
        db.query(models.Payment)
        .filter(
            models.Payment.invoice_id == invoice_id,
            models.Payment.paystack_payment_url.isnot(None),
            models.Payment.status == models.PaymentStatus.pending,
        )
        .order_by(models.Payment.created_at.desc())
        .first()
    )
    paystack_url = paystack_payment.paystack_payment_url if paystack_payment else None
    pdf_bytes = gen_pdf(inv, bank_accounts=bank_accounts, paystack_url=paystack_url)
    return StreamingResponse(
        BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{inv.invoice_number}.pdf"'},
    )


def download_signed_invoice_pdf(db: Session, invoice_id: int):
    inv = get_invoice(db, invoice_id)
    if not inv.custom_pdf_s3_key:
        raise HTTPException(404, "Signed invoice not uploaded yet")
    pdf_bytes = download_bytes(inv.custom_pdf_s3_key)
    return StreamingResponse(
        BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="{inv.invoice_number}_signed.pdf"'},
    )


def download_invoice_cost_of_sales_pdf(db: Session, invoice_id: int):
    inv = get_invoice(db, invoice_id)
    item_rows = (
        db.query(
            models.InvoiceItem.product_id,
            models.Product.product_name,
            func.sum(models.InvoiceItem.quantity).label("qty"),
            func.sum(models.InvoiceItem.cost_price * models.InvoiceItem.quantity).label("cost"),
            func.sum(models.InvoiceItem.line_total).label("revenue"),
        )
        .join(models.Product, models.Product.id == models.InvoiceItem.product_id)
        .filter(models.InvoiceItem.invoice_id == invoice_id)
        .group_by(models.InvoiceItem.product_id, models.Product.product_name)
        .order_by(func.sum(models.InvoiceItem.cost_price * models.InvoiceItem.quantity).desc())
        .all()
    )

    # Keep values in raw kobo — the PDF generator divides by 100 itself
    total_cost = sum(float(product_row.cost or 0) for product_row in item_rows)
    total_revenue = sum(float(product_row.revenue or 0) for product_row in item_rows)
    gross_profit = total_revenue - total_cost
    gross_margin = round(gross_profit / total_revenue * 100, 2) if total_revenue else 0

    report_data = {
        "meta": {
            "invoice_number": inv.invoice_number,
            "customer_name": inv.customer.customer_name if inv.customer else "",
        },
        "summary": {
            "total_cost": total_cost,
            "total_revenue": total_revenue,
            "gross_profit": gross_profit,
            "gross_margin_pct": gross_margin,
        },
        "by_product": [
            {
                "product_id": product_row.product_id,
                "product_name": product_row.product_name,
                "qty": float(product_row.qty or 0),
                "unit_cost_price": (float(product_row.cost or 0) / float(product_row.qty or 0))
                if float(product_row.qty or 0) > 0
                else 0,
                "cost": float(product_row.cost or 0),
                "revenue": float(product_row.revenue or 0),
                "gross_profit": float(product_row.revenue or 0) - float(product_row.cost or 0),
                "margin_pct": round(
                    (float(product_row.revenue or 0) - float(product_row.cost or 0))
                    / float(product_row.revenue)
                    * 100,
                    2,
                )
                if product_row.revenue
                else 0,
            }
            for product_row in item_rows
        ],
        "by_invoice": [
            {
                "invoice_id": inv.id,
                "invoice_number": inv.invoice_number,
                "invoice_date": str(inv.invoice_date),
                "customer_name": inv.customer.customer_name if inv.customer else "",
                "cost": total_cost,
                "revenue": total_revenue,
                "gross_profit": gross_profit,
                "margin_pct": gross_margin,
            }
        ],
    }
    pdf_bytes = generate_cost_of_sales_pdf(
        report_data,
        title_suffix=f" ({inv.invoice_number})",
        cost_only=True,
    )
    return StreamingResponse(
        BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{inv.invoice_number}_cost_of_sales.pdf"'},
    )


async def upload_invoice_pdf(
    db: Session,
    invoice_id: int,
    file: UploadFile,
    additional_emails: Optional[str],
    current_user: models.User,
) -> models.Invoice:
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(400, "Only PDF files are accepted")

    inv = get_invoice(db, invoice_id)
    if inv.custom_pdf_s3_key:
        delete_object(inv.custom_pdf_s3_key)

    content = await file.read()
    s3_key = f"invoices/{invoice_id}/custom_{uuid.uuid4()}.pdf"
    upload_bytes(s3_key, content, "application/pdf")

    inv.custom_pdf_s3_key = s3_key
    db.commit()
    db.refresh(inv)

    recipients: List[str] = [INVOICE_PRIMARY_RECIPIENT]
    if additional_emails:
        recipients.extend([e.strip() for e in additional_emails.split(",") if e.strip()])
    recipients = list(dict.fromkeys(recipients))
    if recipients:
        task = send_invoice_to_recipients_task.delay(inv.id, recipients)
        log_queue_event(
            db,
            task_id=task.id,
            event_type="invoice_email",
            title=f"Send invoice email {inv.invoice_number}",
            requested_by=current_user.id if current_user else None,
            metadata={"invoice_id": inv.id, "recipients": recipients},
        )
    return inv


def remove_invoice_pdf(db: Session, invoice_id: int) -> models.Invoice:
    inv = get_invoice(db, invoice_id)
    if inv.custom_pdf_s3_key:
        delete_object(inv.custom_pdf_s3_key)
        inv.custom_pdf_s3_key = None
        db.commit()
        db.refresh(inv)
    return inv


async def upload_signed_invoice(
    db: Session,
    invoice_id: int,
    file: UploadFile,
    current_user: models.User,
) -> models.Invoice:
    filename = (file.filename or "").lower()
    allowed_image_exts = (".jpg", ".jpeg", ".png", ".webp")
    is_pdf = filename.endswith(".pdf")
    is_image = filename.endswith(allowed_image_exts)
    if not filename or (not is_pdf and not is_image):
        raise HTTPException(400, "Only PDF or image files (JPG, PNG, WEBP) are accepted")

    inv = get_invoice(db, invoice_id)
    if inv.status == models.InvoiceStatus.cancelled:
        raise HTTPException(400, "Cancelled invoice cannot be completed")
    if not inv.quotation_id:
        raise HTTPException(400, "Signed upload is only available for converted invoices")

    if inv.custom_pdf_s3_key:
        delete_object(inv.custom_pdf_s3_key)

    content = await file.read()
    if is_image:
        image = ImageReader(BytesIO(content))
        img_w, img_h = image.getSize()
        page_w, page_h = A4
        scale = min(page_w / img_w, page_h / img_h)
        draw_w = img_w * scale
        draw_h = img_h * scale
        x = (page_w - draw_w) / 2
        y = (page_h - draw_h) / 2
        out = BytesIO()
        pdf = canvas.Canvas(out, pagesize=A4)
        pdf.drawImage(
            image, x, y, width=draw_w, height=draw_h, preserveAspectRatio=True, anchor="c"
        )
        pdf.showPage()
        pdf.save()
        content = out.getvalue()
    s3_key = f"invoices/{invoice_id}/signed_{uuid.uuid4()}.pdf"
    upload_bytes(s3_key, content, "application/pdf")

    inv.custom_pdf_s3_key = s3_key
    inv.status = models.InvoiceStatus.completed
    audit.log(
        db,
        models.AuditAction.update,
        models.AuditEntity.invoice,
        inv.id,
        current_user.id if current_user else None,
        description=f"Uploaded signed invoice for {inv.invoice_number} and marked completed",
        new_values={"status": models.InvoiceStatus.completed.value, "signed_pdf_s3_key": s3_key},
    )
    db.commit()
    db.refresh(inv)

    try:
        send_document_to_make_from_s3(
            doc_type="invoice",
            document_number=inv.invoice_number,
            s3_key=s3_key,
            filename=f"{inv.invoice_number}.pdf",
            customer_name=inv.customer.customer_name if inv.customer else "",
        )
    except Exception:
        pass

    return inv


def generate_invoice_pdf(db: Session, invoice_id: int, current_user: models.User) -> schemas.JobEnqueuedResponse:
    inv = get_invoice(db, invoice_id)
    task = generate_invoice_pdf_task.delay(invoice_id)
    log_queue_event(
        db,
        task_id=task.id,
        event_type="invoice_pdf",
        title=f"Generate invoice PDF {inv.invoice_number}",
        requested_by=current_user.id if current_user else None,
        metadata={"invoice_id": inv.id},
    )
    return schemas.JobEnqueuedResponse(
        task_id=task.id,
        message=f"PDF generation queued for {inv.invoice_number}. Poll /api/v1/jobs/{task.id} for status.",
    )


def send_invoice_email(
    db: Session, invoice_id: int, additional_emails: Optional[List[str]], current_user: models.User
) -> schemas.MessageResponse:
    inv = get_invoice(db, invoice_id)

    emails: List[str] = []
    if inv.customer and inv.customer.email:
        emails.append(inv.customer.email)
    emails.append(INVOICE_PRIMARY_RECIPIENT)
    if additional_emails:
        emails.extend([e.strip() for e in additional_emails if e.strip()])
    emails = list(dict.fromkeys(emails))

    if not emails:
        raise HTTPException(400, "No email addresses to send to")

    task = send_invoice_to_recipients_task.delay(inv.id, emails)
    log_queue_event(
        db,
        task_id=task.id,
        event_type="invoice_email",
        title=f"Send invoice email {inv.invoice_number}",
        requested_by=current_user.id if current_user else None,
        metadata={"invoice_id": inv.id, "recipients": emails},
    )
    return schemas.MessageResponse(message=f"Invoice email queued for {len(emails)} recipient(s)")


def upload_invoice_to_make(
    db: Session, invoice_id: int, additional_emails: Optional[List[str]], current_user: models.User
) -> schemas.MessageResponse:
    inv = get_invoice(db, invoice_id)

    if inv.custom_pdf_s3_key:
        invoice_data = {
            "total_amount": int(round((inv.total_amount or 0) / 100)),
            "amount_paid": int(round((inv.amount_paid or 0) / 100)),
            "delivery_fee": int(round((inv.delivery_fee or 0) / 100)),
            "handling_charge": int(round((inv.handling_charge_amount or 0) / 100)),
            "payment_term": inv.payment_term,
            "delivery_type": inv.delivery_type.value if inv.delivery_type else "",
            "invoice_date": str(inv.invoice_date) if inv.invoice_date else "",
            "due_date": str(inv.due_date) if inv.due_date else "",
            "items": [
                {
                    "product_name": item.product_name or f"Product #{item.product_id}",
                    "quantity": float(item.quantity),
                    "uom": item.uom or "",
                    "unit_price": int(round((item.unit_price or 0) / 100)),
                    "line_total": int(round((item.line_total or 0) / 100)),
                }
                for item in inv.items
            ],
        }
        send_document_to_make_from_s3(
            doc_type="invoice",
            document_number=inv.invoice_number,
            s3_key=inv.custom_pdf_s3_key,
            filename=f"{inv.invoice_number}.pdf",
            customer_name=inv.customer.customer_name if inv.customer else "",
            extra=invoice_data,
        )

    emails: List[str] = [INVOICE_PRIMARY_RECIPIENT]
    if additional_emails:
        emails.extend([e.strip() for e in additional_emails if e.strip()])
    emails = list(dict.fromkeys(emails))
    if not emails:
        raise HTTPException(400, "No email addresses to send to")

    task = send_invoice_to_recipients_task.delay(inv.id, emails)
    log_queue_event(
        db,
        task_id=task.id,
        event_type="invoice_upload_to_make",
        title=f"Upload invoice to make {inv.invoice_number}",
        requested_by=current_user.id if current_user else None,
        metadata={"invoice_id": inv.id, "recipients": emails},
    )
    return schemas.MessageResponse(
        message=f"Invoice upload-to-make queued for {len(emails)} recipient(s)"
    )


def cancel_invoice(db: Session, invoice_id: int, current_user: models.User) -> models.Invoice:
    inv = get_invoice(db, invoice_id)
    if inv.status == models.InvoiceStatus.cancelled:
        raise HTTPException(400, "Invoice is already cancelled")
    if inv.status == models.InvoiceStatus.completed:
        raise HTTPException(400, "Completed invoice cannot be cancelled")
    inv.status = models.InvoiceStatus.cancelled
    audit.log(
        db,
        models.AuditAction.cancel,
        models.AuditEntity.invoice,
        inv.id,
        current_user.id,
        description=f"Cancelled invoice {inv.invoice_number}",
    )
    db.commit()
    db.refresh(inv)
    return inv


def delete_invoice(db: Session, invoice_id: int, current_user: models.User) -> schemas.MessageResponse:
    msg = _delete_invoice_record(db, invoice_id, current_user)
    db.commit()
    return msg
